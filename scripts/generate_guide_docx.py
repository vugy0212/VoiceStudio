import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def create_guide(output_path):
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("VoiceStudio — Korisnički priručnik i upute")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Vodič za kloniranje glasa, dizajn likova, audio knjige, prevođenje i optimizaciju")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    def add_h1(text):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(51, 65, 85)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        return h

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(30, 41, 59)
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.bold = True
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def add_callout(text, title="VAŽAN SAVJET:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(title + " ")
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(14, 116, 144)
        r_b = p.add_run(text)
        r_b.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph()

    # 1. Uvod
    add_h1("1. O aplikaciji VoiceStudio")
    add_p("VoiceStudio je samostalna, lokalna alternativa komercijalnim AI glasovnim servisima (kao što je ElevenLabs). Aplikacija omogućuje precizno kloniranje glasova, kreiranje novih glasovnih identiteta, sinkronizaciju videozapisa i stvaranje audio knjiga. Cijeli sustav radi 100% lokalno na vašem računalu — bez pretplata, bez slanja podataka na internet i bez API ključeva.")

    # 2. Kloniranje glasa
    add_h1("2. Kloniranje glasa (Voice Cloning)")
    add_p("Kloniranje glasa omogućuje vam da iz kratkog uzorka snimke (5 do 15 sekundi) precizno prekopirate boju glasa, intonaciju i način govora bilo koje osobe.")
    add_h2("Koraci za savršeno kloniranje:")
    add_bullet("Snimka mora biti bez pozadinske glazbe, jeke i šumova. Format može biti WAV, MP3 ili M4A.", "1. Priprema audio isječka: ")
    add_bullet("Učitajte audio datoteku u polje 'Reference audio'.", "2. Učitavanje u VoiceStudio: ")
    add_bullet("U glavno tekstualno polje upišite rečenice koje želite da klonirani glas izgovori.", "3. Unos teksta: ")
    add_bullet("Ispod učitanog audio isječka kliknite na malu strelicu 'Optional details'. U polje 'Transcript' upišite točan tekst koji se izgovara u toj referentnoj snimci. To drastično ubrzava proces.", "4. Unos transkripta referentnog zvuka: ")
    add_bullet("Kliknite ružičasti gumb 'Synthesize' na dnu ekrana. Za nekoliko sekundi vaš audio će biti spreman za preslušavanje i preuzimanje.", "5. Generiranje: ")

    add_callout(
        "Ako polje 'Transcript' u 'Optional details' ostavite prazno, aplikacija mora najprije pokrenuti Whisper AI model da sama presluša i prepozna riječi snimke. Unosom teksta štedite vrijeme i osiguravate trenutno generiranje!",
        "KLJUČNA UPUTA ZA MAKSIMALNU BRZINU:"
    )

    # 3. Kontrola brzine i parametri
    add_h1("3. Kontrola brzine i postavki zvuka (Production Overrides)")
    add_p("Točno iznad gumba 'Synthesize' nalazi se izbornik 'Production Overrides ^'. Klikom na njega otvaraju se precizne kontrole za glas:")
    add_bullet("Zadana vrijednost je 1.0x. Ako je govor prebrz, smanjite klizač na 0.85x ili 0.9x. Ako želite brži govor, povisite na 1.1x ili više.", "Speed (Brzina izgovora): ")
    add_bullet("Određuje koliko se model strogo drži referentnog glasa. Zadana vrijednost (2.0) je optimalna za većinu glasova.", "CFG / Guidance Scale: ")
    add_bullet("Broj koraka obrade difuzije (zadano 16 do 32). Više koraka daje nešto čišći zvuk, dok manje koraka omogućuje brže generiranje.", "Steps (Kvaliteta uzorkovanja): ")
    add_bullet("Uklanja tihi šum iz pozadine sintetičkog zvuka.", "Denoise (Čišćenje šuma): ")
    add_bullet("Možete upisati točan broj sekundi koliko želite da rečenica traje ako trebate savršeno uklapanje u video.", "Duration (Fiksno trajanje): ")

    # 4. Interpunkcija i emocije
    add_h1("4. Umetanje pauza i ritam govora")
    add_p("AI modeli za govor iznimno paze na interpunkciju:")
    add_bullet("Zarez stvara kratku, prirodnu stanku za udah.", "Zarez (,): ")
    add_bullet("Trotočje stvara dužu, dramatičnu pauzu.", "Trotočje (...): ")
    add_bullet("Duga crtica stvara jasnu promjenu ritma rečenice.", "Crtica (—): ")
    add_bullet("Upitnik i uskličnik prirodno podižu visinu tona na kraju rečenice.", "Upitnik (?) i Uskličnik (!): ")

    # 5. Ostale funkcionalnosti
    add_h1("5. Ostale mogućnosti VoiceStudia")
    add_bullet("Omogućuje generiranje potpuno novog, nepostojećeg glasa pomoću opisa (npr. 'stariji duboki muški glas s promuklim tonom').", "Voice Design (Dizajn glasa): ")
    add_bullet("Automatsko prevođenje i presnimavanje videozapisa uz očuvanje originalnih zvukova u pozadini.", "Video Dubbing (Sinkronizacija): ")
    add_bullet("Uređivač za knjige i priče u kojem možete dodijeliti različite glasove različitim likovima (npr. Narator, Lik 1, Lik 2) i izvesti cijelo poglavlje u MP3.", "Audiobook & Stories (Audio knjige): ")
    add_bullet("Mogućnost diktiranja govora izravno u bilo koji program na računalu.", "Dictation (Diktiranje): ")

    # 6. Optimizacija i rješavanje problema
    add_h1("6. Savjeti za rad na vašem računalu (RTX 4060)")
    add_bullet("VoiceStudio koristi vašu NVIDIA grafičku karticu preko CUDA ubrzanja. Na 10 sekundi teksta potrebno je samo 1.5 do 3.5 sekundi obrade.", "Izvrsne performanse: ")
    add_bullet("Ako koristite LM Studio ili druge lokalne AI programe, ugasite ih dok radite u VoiceStudiju kako bi svih 8 GB video memorije bilo dostupno za generiranje glasa.", "Gašenje drugih AI alata: ")
    add_bullet("Na Radnoj površini kreiran je prečac 'Pokreni VoiceStudio' koji jednim klikom rješava portove i otvara aplikaciju u pregledniku.", "Brzo pokretanje: ")

    doc.save(output_path)

if __name__ == "__main__":
    output_docx = r"c:\Users\PC\Desktop\VoiceStudio_Korisnicki_Vodic.docx"
    create_guide(output_docx)
    print(f"Docx generated at {output_docx}")
